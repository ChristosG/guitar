import io

from docx import Document
from fastapi.testclient import TestClient

from app.db import SessionLocal
from app.main import app
from app.models.block import Block

client = TestClient(app)


def _mk_tree(db):
    course = Block(kind="course", title="Ήχος και Ενισχυτές", is_template=True,
                   order=0, plane="content", target_profile={"level": "μεσαίο"})
    db.add(course); db.flush()
    m1 = Block(kind="module", title="Βασικές αρχές ήχου", parent_id=course.id, order=0, plane="content")
    db.add(m1); db.flush()
    l1 = Block(kind="lesson", title="Τι είναι το gain", parent_id=m1.id, order=0,
               plane="content", est_minutes=55,
               meta={"draft_status": "ready",
                     "citations": [{"source_id": "s1", "source_ref": "b1",
                                    "source_title": "Getting Great Guitar Sounds", "page": 19}]})
    db.add(l1); db.flush()
    db.add(Block(kind="segment", title="Θεωρία", parent_id=l1.id, order=0, plane="content",
                 body="Το gain καθορίζει την προενίσχυση του σήματος.",
                 meta={"section": "theory"}))
    db.add(Block(kind="segment", title="Ασκήσεις", parent_id=l1.id, order=1, plane="content",
                 body="Άσκηση 1: σύγκρινε clean και overdriven ήχο.",
                 meta={"section": "exercises"}))
    l2 = Block(kind="lesson", title="Ακόμα γράφεται", parent_id=m1.id, order=1,
               plane="content", meta={"draft_status": "queued"})
    db.add(l2); db.commit()
    return course


def test_export_docx_structure_and_greek_content():
    db = SessionLocal()
    try:
        course = _mk_tree(db)
        r = client.get(f"/curricula/{course.id}/export.docx")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert "attachment" in r.headers["content-disposition"]

        doc = Document(io.BytesIO(r.content))
        headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
        texts = [p.text for p in doc.paragraphs]

        assert ("Heading 1", "Βασικές αρχές ήχου") in headings
        assert any(s == "Heading 2" and "Τι είναι το gain" in t for s, t in headings)
        assert ("Heading 3", "Θεωρία") in headings
        assert ("Heading 3", "Ασκήσεις") in headings
        assert "Το gain καθορίζει την προενίσχυση του σήματος." in texts
        # heading ORDER: module before its lesson before its sections
        h_texts = [t for _, t in headings]
        assert h_texts.index("Βασικές αρχές ήχου") < h_texts.index(next(t for t in h_texts if "Τι είναι το gain" in t))

        # spec §4: NO sources anywhere in the document
        joined = "\n".join(texts)
        assert "Getting Great Guitar Sounds" not in joined
        assert "Πηγές" not in joined

        # undrafted lesson: present with a placeholder, not silently missing
        assert any("Ακόμα γράφεται" in t for t in h_texts)
    finally:
        db.close()


def test_export_docx_marks_failed_and_drafting_lessons_with_stale_segments():
    """A lesson that has OLD segments (from a prior successful draft) but is
    now `failed` or `drafting` a redraft must still render the undrafted
    note — after its stale content, not instead of it — so the export never
    reads as a clean finished lesson when the board itself shows a spinner
    or a retry button (review fix: this used to only fire for zero-segment
    lessons)."""
    db = SessionLocal()
    try:
        course = _mk_tree(db)
        m1 = next(b for b in course.children if b.kind == "module")

        failed = Block(kind="lesson", title="Παλιό μάθημα, απέτυχε το ξαναγράψιμο",
                       parent_id=m1.id, order=2, plane="content",
                       meta={"draft_status": "failed", "error": "boom"})
        db.add(failed); db.flush()
        db.add(Block(kind="segment", title="Θεωρία", parent_id=failed.id, order=0,
                     plane="content", body="Παλιό, ίσως ξεπερασμένο κείμενο."))

        drafting = Block(kind="lesson", title="Ξαναγράφεται τώρα",
                         parent_id=m1.id, order=3, plane="content",
                         meta={"draft_status": "drafting"})
        db.add(drafting); db.flush()
        db.add(Block(kind="segment", title="Θεωρία", parent_id=drafting.id, order=0,
                     plane="content", body="Ακόμα παλιότερο κείμενο."))
        db.commit()

        r = client.get(f"/curricula/{course.id}/export.docx")
        assert r.status_code == 200
        doc = Document(io.BytesIO(r.content))
        texts = [p.text for p in doc.paragraphs]
        headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]

        # the stale prose is still there (not silently dropped)...
        assert "Παλιό, ίσως ξεπερασμένο κείμενο." in texts
        assert "Ακόμα παλιότερο κείμενο." in texts
        # ...but the undrafted note follows it for BOTH lessons — plus the
        # THIRD occurrence `_mk_tree`'s own zero-segment "queued" lesson (l2)
        # already contributes (the pre-existing zero-segments case, untouched
        # by this fix).
        undrafted_note = "— Το μάθημα δεν έχει συνταχθεί ακόμα. —"
        assert texts.count(undrafted_note) == 3

        stale_idx = texts.index("Παλιό, ίσως ξεπερασμένο κείμενο.")
        note_indices = [i for i, t in enumerate(texts) if t == undrafted_note]
        assert any(i > stale_idx for i in note_indices)

        # the queued lesson (zero segments, unrelated to this fix) still gets
        # exactly the one note it always did.
        assert any("Ακόμα γράφεται" in t for _, t in headings)
    finally:
        db.close()


def test_export_docx_404_for_non_roots():
    db = SessionLocal()
    try:
        course = _mk_tree(db)
        module_id = next(b.id for b in course.children)
        assert client.get(f"/curricula/{module_id}/export.docx").status_code == 404
    finally:
        db.close()


def test_export_docx_renders_the_inline_marks_as_real_word_runs():
    """`**bold**`, `*italic*` and `<u>underline</u>` are formatting, not text:
    the Word file must carry them as run properties. A tutor who prints this
    for a student must not be handed a page full of asterisks."""
    db = SessionLocal()
    try:
        course = _mk_tree(db)
        lesson = next(
            b for m in course.children for b in m.children
            if b.kind == "lesson" and "gain" in b.title
        )
        db.add(Block(kind="segment", title="Πηνία", parent_id=lesson.id, order=2,
                     plane="content", meta={"section": "theory"},
                     body="Το **humbucker** έχει *δύο* πηνία <u>σε σειρά</u>"))
        db.commit()

        r = client.get(f"/curricula/{course.id}/export.docx")
        assert r.status_code == 200
        doc = Document(io.BytesIO(r.content))

        para = next(p for p in doc.paragraphs if "humbucker" in p.text)
        # the markers themselves are gone from the visible text...
        assert para.text == "Το humbucker έχει δύο πηνία σε σειρά"
        runs = para.runs
        assert [r.text for r in runs if r.bold] == ["humbucker"]
        assert [r.text for r in runs if r.italic] == ["δύο"]
        assert [r.text for r in runs if r.underline] == ["σε σειρά"]
        # ...and no run anywhere in the document still shows a marker.
        every_run = [run.text for p in doc.paragraphs for run in p.runs]
        assert not any("*" in t or "<u>" in t or "</u>" in t for t in every_run)
        # the unmarked words are plain, not accidentally inheriting the span
        plain = [r.text for r in runs if not (r.bold or r.italic or r.underline)]
        assert plain == ["Το ", " έχει ", " πηνία "]
    finally:
        db.close()
